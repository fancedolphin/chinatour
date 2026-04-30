#!/usr/bin/env python3

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import uuid
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


PROJECT_URL = "https://ogodnvjaiwelqmjqkvda.supabase.co"
DEFAULT_DIANPING_CSV = "/mnt/d/dianping_data/dianping.csv"
DEFAULT_DISH_CSV = "/mnt/d/dianping_data/dish.csv"
DEFAULT_POSITION_XLSX = "/mnt/d/dianping_data/supabase_database_full.xlsx"
DEFAULT_SQL_OUTPUT = "/tmp/import_dianping_restaurants.sql"
DEFAULT_FUNCTION_SLUG = "import-restaurant-batch"

UUID_NAMESPACE = uuid.UUID("52f2fefe-7b7f-4ae5-b4d9-c30f4466acdc")

CITY_PATTERNS = [
    "北京",
    "上海",
    "佛山",
    "顺德",
    "太原",
    "桃源县",
    "桃源",
    "常德",
    "安顺",
    "齐齐哈尔",
    "钦州",
    "西虹桥",
]

CUISINE_RULES = [
    ("面馆", "面食"),
    ("烧鸡", "烧鸡"),
    ("烧鹅", "烧味"),
    ("小炒", "小炒"),
    ("私房菜", "私房菜"),
    ("擂茶", "客家菜"),
    ("饭店", "地方菜"),
    ("餐馆", "家常菜"),
    ("酒店", "地方菜"),
]

KEYWORD_RULES = [
    ("排队", ["排队", "等座", "翻台", "排号"]),
    ("停车", ["停车", "车位", "打车"]),
    ("环境", ["环境", "装修", "包间", "大排档"]),
    ("服务", ["服务", "服务员", "态度"]),
    ("性价比", ["性价比", "便宜", "划算", "价格不贵"]),
    ("游客多", ["游客", "外地朋友", "外地人"]),
]


def normalize_name(value: str) -> str:
    text = (value or "").strip()
    text = re.sub(r"[·•\s]", "", text)
    text = re.sub(r"（[^）]*）|\([^)]*\)", "", text)
    return text.lower()


def collapse_ws(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def make_uuid(label: str) -> str:
    return str(uuid.uuid5(UUID_NAMESPACE, label))


def sql_string(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + value.replace("\\", "\\\\").replace("'", "''") + "'"


def sql_array(values: list[str] | None) -> str:
    if not values:
        return "NULL"
    escaped = [sql_string(collapse_ws(value)) for value in values if collapse_ws(value)]
    if not escaped:
        return "NULL"
    return "ARRAY[" + ", ".join(escaped) + "]"


def sql_vector(values: list[float] | None) -> str:
    if not values:
        return "NULL"
    payload = "[" + ",".join(f"{value:.9f}".rstrip("0").rstrip(".") for value in values) + "]"
    return sql_string(payload) + "::vector"


def vector_literal(values: list[float] | None) -> str | None:
    if not values:
        return None
    return "[" + ",".join(f"{value:.9f}".rstrip("0").rstrip(".") for value in values) + "]"


def deterministic_hash(parts: list[str]) -> str:
    joined = "|".join(parts)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


def load_position_rows(path: Path) -> tuple[dict[str, dict], dict[str, dict]]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook["Position"]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index = {name: position for position, name in enumerate(headers)}

    exact: dict[str, dict] = {}
    normalized: dict[str, dict] = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = row[index["name"]]
        if not name:
            continue
        record = {
            "name": str(name).strip(),
            "province": row[index["province"]],
            "city": row[index["city"]],
            "district": row[index["district"]],
            "address": row[index["address"]],
            "longitude": row[index["longitude"]],
            "latitude": row[index["latitude"]],
            "link_dianping": row[index["link_dianping"]],
        }
        exact[record["name"]] = record
        normalized.setdefault(normalize_name(record["name"]), record)

    return exact, normalized


def load_reviews(path: Path) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    seen_hashes: set[str] = set()

    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            name = collapse_ws(row.get("餐馆名称", ""))
            content = collapse_ws(row.get("评论内容", ""))
            if not name or not content:
                continue

            fingerprint = deterministic_hash(
                [
                    name,
                    collapse_ws(row.get("用户名", "")),
                    collapse_ws(row.get("评分", "")),
                    collapse_ws(row.get("评论时间", "")),
                    content,
                ]
            )
            if fingerprint in seen_hashes:
                continue
            seen_hashes.add(fingerprint)

            grouped[name].append(
                {
                    "user": collapse_ws(row.get("用户名", "")),
                    "rating": collapse_ws(row.get("评分", "")),
                    "content": content,
                    "time": collapse_ws(row.get("评论时间", "")),
                }
            )

    return grouped


def build_alias_map(review_names: set[str], dish_names: set[str]) -> dict[str, str]:
    alias_map: dict[str, str] = {}

    for name in sorted(dish_names):
        if name in review_names:
            alias_map[name] = name
            continue

        candidates = [review_name for review_name in review_names if name and name in review_name]
        if len(candidates) == 1:
            alias_map[name] = candidates[0]
            continue

        alias_map[name] = name

    return alias_map


def load_dishes(path: Path, alias_map: dict[str, str]) -> dict[str, list[dict]]:
    grouped: dict[str, dict[str, dict]] = defaultdict(dict)

    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            raw_name = collapse_ws(row.get("restaurantName", ""))
            dish_name = collapse_ws(row.get("dishName", ""))
            if not raw_name or not dish_name:
                continue

            canonical_name = alias_map.get(raw_name, raw_name)
            existing = grouped[canonical_name].get(dish_name)
            recommend_count = int(row.get("recommendCount", "0") or 0)
            image_url = collapse_ws(row.get("picUrl", "")) or None

            if existing:
                existing["recommend_count"] += recommend_count
                existing["image_url"] = existing["image_url"] or image_url
                continue

            grouped[canonical_name][dish_name] = {
                "dish_name": dish_name,
                "recommend_count": recommend_count,
                "image_url": image_url,
            }

    return {name: list(dishes.values()) for name, dishes in grouped.items()}


def infer_city(name: str, reviews: list[dict], matched_position: dict | None) -> str | None:
    if matched_position and matched_position.get("city"):
        return collapse_ws(str(matched_position["city"]))

    content_blob = " ".join(review["content"] for review in reviews)
    counts = Counter()
    for city in CITY_PATTERNS:
        if city and city in content_blob:
            counts[city] += content_blob.count(city)

    if counts:
        city = counts.most_common(1)[0][0]
        if city == "桃源":
            return "桃源县"
        if city == "顺德":
            return "佛山"
        if city == "西虹桥":
            return "上海"
        return city

    return None


def infer_cuisine(name: str, top_dishes: list[str]) -> str | None:
    for keyword, cuisine in CUISINE_RULES:
        if keyword in name:
            return cuisine

    joined = " ".join(top_dishes)
    if "面" in joined:
        return "面食"
    if "鸡" in joined:
        return "地方菜"
    if "烧鹅" in joined:
        return "烧味"
    return None


def parse_price_range(reviews: list[dict]) -> str | None:
    values: list[int] = []
    pattern = re.compile(r"人均\s*[:：]?\s*(\d{1,4})")
    for review in reviews:
        for match in pattern.findall(review["content"]):
            value = int(match)
            if 1 <= value <= 5000:
                values.append(value)

    if not values:
        return None

    low = min(values)
    high = max(values)
    if low == high:
        return f"约¥{low}/人"
    return f"约¥{low}-{high}/人"


def build_keywords(reviews: list[dict]) -> list[str]:
    counts = Counter()
    for review in reviews:
        content = review["content"]
        for label, rules in KEYWORD_RULES:
            if any(rule in content for rule in rules):
                counts[label] += 1

    return [label for label, _ in counts.most_common(3)]


def build_description(
    name: str,
    city: str | None,
    avg_rating: float | None,
    review_count: int,
    top_dishes: list[str],
    keywords: list[str],
) -> str:
    parts = [f"综合{review_count}条大众点评评论整理"]
    if city:
        parts.append(f"位于{city}")
    if avg_rating is not None:
        parts.append(f"平均评分{avg_rating:.1f}")
    if top_dishes:
        parts.append("推荐菜包括" + "、".join(top_dishes[:4]))
    if keywords:
        parts.append("高频体验关键词：" + "、".join(keywords))
    return "；".join(parts) + "。"


def compute_popularity(review_count: int, total_recommends: int) -> float:
    raw = 0.30 + min(0.65, (math.log(review_count + 1) + math.log(total_recommends + 1)) / 10)
    return round(min(0.95, raw), 2)


def generate_embedding(project_url: str, anon_key: str, text: str) -> list[float]:
    payload = json.dumps({"text": text}).encode("utf-8")
    request = urllib.request.Request(
        f"{project_url}/functions/v1/generate-embedding",
        data=payload,
        headers={
            "Content-Type": "application/json",
            "apikey": anon_key,
            "Authorization": f"Bearer {anon_key}",
        },
        method="POST",
    )

    with urllib.request.urlopen(request, timeout=60) as response:
        body = json.loads(response.read().decode("utf-8"))

    embedding = body.get("embedding")
    if not isinstance(embedding, list) or len(embedding) != 768:
        raise ValueError("embedding response is invalid")
    return embedding


def post_json(url: str, anon_key: str, payload: dict) -> dict:
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=data,
        headers={
            "Content-Type": "application/json",
            "apikey": anon_key,
            "Authorization": f"Bearer {anon_key}",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} calling {url}: {details}") from exc


def build_restaurant_payloads(
    reviews_by_name: dict[str, list[dict]],
    dishes_by_name: dict[str, list[dict]],
    position_exact: dict[str, dict],
    position_normalized: dict[str, dict],
) -> tuple[list[dict], list[dict], list[dict]]:
    canonical_names = sorted(set(reviews_by_name) | set(dishes_by_name))
    restaurants: list[dict] = []
    destinations: dict[str, dict] = {}
    dishes: list[dict] = []

    for name in canonical_names:
        reviews = reviews_by_name.get(name, [])
        dish_rows = sorted(
            dishes_by_name.get(name, []),
            key=lambda item: (-item["recommend_count"], item["dish_name"]),
        )

        matched_position = position_exact.get(name) or position_normalized.get(normalize_name(name))
        city = infer_city(name, reviews, matched_position)
        destination_id = None
        if city:
            destination_id = make_uuid(f"destination:{city}")
            destinations[city] = {
                "id": destination_id,
                "name": city,
                "country": "中国",
                "description": f"{city}餐饮目的地知识库（由大众点评评论与推荐菜导入）。",
                "best_season": None,
                "average_budget_daily": None,
                "currency": "CNY",
                "timezone": "Asia/Shanghai",
            }

        ratings = [float(review["rating"]) for review in reviews if review["rating"]]
        avg_rating = round(sum(ratings) / len(ratings), 1) if ratings else None
        review_count = len(reviews)
        total_recommends = sum(dish["recommend_count"] for dish in dish_rows)
        top_dishes = [dish["dish_name"] for dish in dish_rows[:5]]
        keywords = build_keywords(reviews)
        description = build_description(name, city, avg_rating, review_count, top_dishes, keywords)
        cuisine_type = infer_cuisine(name, top_dishes)
        price_range = parse_price_range(reviews)
        menu_image_url = dish_rows[0]["image_url"] if dish_rows else None
        restaurant_id = make_uuid(f"restaurant:{city or 'unknown'}:{name}")

        restaurants.append(
            {
                "id": restaurant_id,
                "destination_id": destination_id,
                "name": name,
                "name_en": None,
                "cuisine_type": cuisine_type,
                "description": description,
                "location_lat": matched_position.get("latitude") if matched_position else None,
                "location_lng": matched_position.get("longitude") if matched_position else None,
                "address": matched_position.get("address") if matched_position else None,
                "price_range": price_range,
                "meal_type": None,
                "specialties": top_dishes or None,
                "dietary_options": None,
                "opening_hours": None,
                "reservation_required": any("预约" in review["content"] for review in reviews),
                "popularity_score": compute_popularity(review_count, total_recommends),
                "avg_rating": avg_rating,
                "review_count": review_count,
                "menu_image_url": menu_image_url,
                "city": city,
                "dianping_url": collapse_ws(str(matched_position["link_dianping"]))
                    if matched_position and matched_position.get("link_dianping")
                    else None,
            }
        )

        for order_index, dish in enumerate(dish_rows):
            dishes.append(
                {
                    "id": make_uuid(f"dish:{restaurant_id}:{dish['dish_name']}"),
                    "restaurant_id": restaurant_id,
                    "name": dish["dish_name"],
                    "name_en": None,
                    "description": f"大众点评推荐 {dish['recommend_count']} 次",
                    "image_url": dish["image_url"],
                    "allergens": [],
                    "recommend_count": dish["recommend_count"],
                    "order_index": order_index,
                }
            )

    return list(destinations.values()), restaurants, dishes


def attach_embeddings(project_url: str, anon_key: str, destinations: list[dict], restaurants: list[dict]) -> None:
    for destination in destinations:
        text = " ".join(
            part
            for part in [
                destination["name"],
                destination["description"],
                destination["country"],
            ]
            if part
        )
        destination["embedding"] = generate_embedding(project_url, anon_key, text)

    for restaurant in restaurants:
        text = " ".join(
            part
            for part in [
                restaurant["city"] or "",
                restaurant["name"],
                restaurant["cuisine_type"] or "",
                restaurant["description"],
                " ".join(restaurant["specialties"] or []),
                restaurant["address"] or "",
            ]
            if part
        )
        restaurant["embedding"] = generate_embedding(project_url, anon_key, text)


def upload_batches(
    project_url: str,
    anon_key: str,
    function_slug: str,
    destinations: list[dict],
    restaurants: list[dict],
    dishes: list[dict],
    batch_size: int,
) -> list[dict]:
    results: list[dict] = []
    endpoint = f"{project_url}/functions/v1/{function_slug}"
    restaurant_ids = {restaurant["id"] for restaurant in restaurants}
    dishes_by_restaurant: dict[str, list[dict]] = defaultdict(list)

    for dish in dishes:
        if dish["restaurant_id"] in restaurant_ids:
            dishes_by_restaurant[dish["restaurant_id"]].append(
                {
                    **dish,
                    "allergens": dish.get("allergens") or [],
                }
            )

    destination_payload = [
        {
            **destination,
            "embedding": vector_literal(destination.get("embedding")),
        }
        for destination in destinations
    ]

    if destination_payload:
        results.append(
            post_json(
                endpoint,
                anon_key,
                {
                    "destinations": destination_payload,
                    "restaurants": [],
                    "dishes": [],
                },
            )
        )

    for start in range(0, len(restaurants), batch_size):
        restaurant_batch = restaurants[start : start + batch_size]
        restaurant_payload = [
            {
                **restaurant,
                "embedding": vector_literal(restaurant.get("embedding")),
            }
            for restaurant in restaurant_batch
        ]
        dish_payload: list[dict] = []
        for restaurant in restaurant_batch:
            dish_payload.extend(dishes_by_restaurant.get(restaurant["id"], []))

        results.append(
            post_json(
                endpoint,
                anon_key,
                {
                    "destinations": [],
                    "restaurants": restaurant_payload,
                    "dishes": dish_payload,
                },
            )
        )

    return results


def build_sql(destinations: list[dict], restaurants: list[dict], dishes: list[dict]) -> str:
    lines = ["BEGIN;"]

    for destination in destinations:
        lines.append(
            """
INSERT INTO public.destinations (
  id, name, country, description, best_season, average_budget_daily, currency, timezone, embedding, updated_at
) VALUES (
  {id}, {name}, {country}, {description}, {best_season}, {average_budget_daily}, {currency}, {timezone}, {embedding}, NOW()
)
ON CONFLICT (id) DO UPDATE SET
  name = EXCLUDED.name,
  country = EXCLUDED.country,
  description = EXCLUDED.description,
  best_season = EXCLUDED.best_season,
  average_budget_daily = EXCLUDED.average_budget_daily,
  currency = EXCLUDED.currency,
  timezone = EXCLUDED.timezone,
  embedding = EXCLUDED.embedding,
  updated_at = NOW();
""".format(
                id=sql_string(destination["id"]),
                name=sql_string(destination["name"]),
                country=sql_string(destination["country"]),
                description=sql_string(destination["description"]),
                best_season=sql_string(destination["best_season"]),
                average_budget_daily=sql_string(destination["average_budget_daily"]),
                currency=sql_string(destination["currency"]),
                timezone=sql_string(destination["timezone"]),
                embedding=sql_vector(destination.get("embedding")),
            ).strip()
        )

    for restaurant in restaurants:
        lines.append(
            """
INSERT INTO public.restaurants (
  id, destination_id, name, name_en, cuisine_type, description, location_lat, location_lng, address,
  price_range, meal_type, specialties, dietary_options, opening_hours, reservation_required,
  popularity_score, avg_rating, review_count, menu_image_url, dianping_url, embedding, updated_at
) VALUES (
  {id}, {destination_id}, {name}, {name_en}, {cuisine_type}, {description}, {location_lat}, {location_lng}, {address},
  {price_range}, {meal_type}, {specialties}, {dietary_options}, {opening_hours}, {reservation_required},
  {popularity_score}, {avg_rating}, {review_count}, {menu_image_url}, {dianping_url}, {embedding}, NOW()
)
ON CONFLICT (id) DO UPDATE SET
  destination_id = EXCLUDED.destination_id,
  name = EXCLUDED.name,
  name_en = EXCLUDED.name_en,
  cuisine_type = EXCLUDED.cuisine_type,
  description = EXCLUDED.description,
  location_lat = EXCLUDED.location_lat,
  location_lng = EXCLUDED.location_lng,
  address = EXCLUDED.address,
  price_range = EXCLUDED.price_range,
  meal_type = EXCLUDED.meal_type,
  specialties = EXCLUDED.specialties,
  dietary_options = EXCLUDED.dietary_options,
  opening_hours = EXCLUDED.opening_hours,
  reservation_required = EXCLUDED.reservation_required,
  popularity_score = EXCLUDED.popularity_score,
  avg_rating = EXCLUDED.avg_rating,
  review_count = EXCLUDED.review_count,
  menu_image_url = EXCLUDED.menu_image_url,
  dianping_url = EXCLUDED.dianping_url,
  embedding = EXCLUDED.embedding,
  updated_at = NOW();
""".format(
                id=sql_string(restaurant["id"]),
                destination_id=sql_string(restaurant["destination_id"]),
                name=sql_string(restaurant["name"]),
                name_en=sql_string(restaurant["name_en"]),
                cuisine_type=sql_string(restaurant["cuisine_type"]),
                description=sql_string(restaurant["description"]),
                location_lat=str(restaurant["location_lat"]) if restaurant["location_lat"] is not None else "NULL",
                location_lng=str(restaurant["location_lng"]) if restaurant["location_lng"] is not None else "NULL",
                address=sql_string(restaurant["address"]),
                price_range=sql_string(restaurant["price_range"]),
                meal_type=sql_string(restaurant["meal_type"]),
                specialties=sql_array(restaurant["specialties"]),
                dietary_options=sql_array(restaurant["dietary_options"]),
                opening_hours="NULL",
                reservation_required="TRUE" if restaurant["reservation_required"] else "FALSE",
                popularity_score=str(restaurant["popularity_score"]),
                avg_rating=str(restaurant["avg_rating"]) if restaurant["avg_rating"] is not None else "NULL",
                review_count=str(restaurant["review_count"]),
                menu_image_url=sql_string(restaurant["menu_image_url"]),
                dianping_url=sql_string(restaurant.get("dianping_url")),
                embedding=sql_vector(restaurant.get("embedding")),
            ).strip()
        )

    restaurant_ids = [sql_string(restaurant["id"]) for restaurant in restaurants]
    if restaurant_ids:
        lines.append(
            "DELETE FROM public.restaurant_dishes WHERE restaurant_id = ANY (ARRAY[{restaurant_ids}]::uuid[]);".format(
                restaurant_ids=", ".join(restaurant_ids)
            )
        )

    for dish in dishes:
        lines.append(
            """
INSERT INTO public.restaurant_dishes (
  id, restaurant_id, name, name_en, description, image_url, allergens, recommend_count, order_index, updated_at
) VALUES (
  {id}, {restaurant_id}, {name}, {name_en}, {description}, {image_url}, {allergens}, {recommend_count}, {order_index}, NOW()
)
ON CONFLICT (id) DO UPDATE SET
  restaurant_id = EXCLUDED.restaurant_id,
  name = EXCLUDED.name,
  name_en = EXCLUDED.name_en,
  description = EXCLUDED.description,
  image_url = EXCLUDED.image_url,
  allergens = EXCLUDED.allergens,
  recommend_count = EXCLUDED.recommend_count,
  order_index = EXCLUDED.order_index,
  updated_at = NOW();
""".format(
                id=sql_string(dish["id"]),
                restaurant_id=sql_string(dish["restaurant_id"]),
                name=sql_string(dish["name"]),
                name_en=sql_string(dish["name_en"]),
                description=sql_string(dish["description"]),
                image_url=sql_string(dish["image_url"]),
                allergens=sql_array(dish["allergens"]),
                recommend_count=str(dish["recommend_count"]),
                order_index=str(dish["order_index"]),
            ).strip()
        )

    lines.append("COMMIT;")
    return "\n\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="Build restaurant import SQL from dianping and dish csv files.")
    parser.add_argument("--dianping-csv", default=DEFAULT_DIANPING_CSV)
    parser.add_argument("--dish-csv", default=DEFAULT_DISH_CSV)
    parser.add_argument("--position-xlsx", default=DEFAULT_POSITION_XLSX)
    parser.add_argument("--project-url", default=PROJECT_URL)
    parser.add_argument("--anon-key", default="")
    parser.add_argument("--sql-output", default=DEFAULT_SQL_OUTPUT)
    parser.add_argument("--skip-embedding", action="store_true")
    parser.add_argument("--apply-via-function", action="store_true")
    parser.add_argument("--function-slug", default=DEFAULT_FUNCTION_SLUG)
    parser.add_argument("--batch-size", type=int, default=10)
    args = parser.parse_args()

    dianping_path = Path(args.dianping_csv)
    dish_path = Path(args.dish_csv)
    position_path = Path(args.position_xlsx)
    sql_output = Path(args.sql_output)

    reviews_by_name = load_reviews(dianping_path)
    review_names = set(reviews_by_name)
    dish_names = {
        collapse_ws(row["restaurantName"])
        for row in csv.DictReader(dish_path.open("r", encoding="utf-8-sig", newline=""))
        if row.get("restaurantName")
    }
    alias_map = build_alias_map(review_names, dish_names)
    dishes_by_name = load_dishes(dish_path, alias_map)
    position_exact, position_normalized = load_position_rows(position_path)

    destinations, restaurants, dishes = build_restaurant_payloads(
        reviews_by_name, dishes_by_name, position_exact, position_normalized
    )

    if not args.skip_embedding:
        if not args.anon_key:
            parser.error("--anon-key is required unless --skip-embedding is used")
        attach_embeddings(args.project_url, args.anon_key, destinations, restaurants)

    sql_output.write_text(build_sql(destinations, restaurants, dishes), encoding="utf-8")

    upload_results = []
    if args.apply_via_function:
        if not args.anon_key:
            parser.error("--anon-key is required when --apply-via-function is used")
        upload_results = upload_batches(
            args.project_url,
            args.anon_key,
            args.function_slug,
            destinations,
            restaurants,
            dishes,
            max(1, args.batch_size),
        )

    summary = {
        "destinations": len(destinations),
        "restaurants": len(restaurants),
        "restaurant_dishes": len(dishes),
        "with_position_match": sum(
            1 for restaurant in restaurants if restaurant["address"] or restaurant["location_lat"] is not None
        ),
        "without_destination": [restaurant["name"] for restaurant in restaurants if not restaurant["destination_id"]],
        "sql_output": str(sql_output),
        "upload_batches": len(upload_results),
        "upload_results": upload_results,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
