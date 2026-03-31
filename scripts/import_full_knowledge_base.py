#!/usr/bin/env python3

from __future__ import annotations

import argparse
import ast
import json
import math
import re
import sys
import time
import urllib.error
import urllib.request
import uuid
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


PROJECT_URL = "https://ogodnvjaiwelqmjqkvda.supabase.co"
DEFAULT_POSITION_XLSX = "/mnt/d/dianping_data/supabase_database_full.xlsx"
DEFAULT_NOTES_DIR = "/mnt/d/dianping_data/datas/excel_datas"
DEFAULT_FUNCTION_SLUG = "import-knowledge-batch"
UUID_NAMESPACE = uuid.UUID("52f2fefe-7b7f-4ae5-b4d9-c30f4466acdc")

SKIP_NOTE_FILES = {"restaurant_notes_summary.xlsx"}

CITY_ALIASES = {
    "北京": "北京市",
    "北京市": "北京市",
    "上海": "上海市",
    "上海市": "上海市",
    "天津": "天津市",
    "天津市": "天津市",
    "重庆": "重庆市",
    "重庆市": "重庆市",
    "佛山": "佛山市",
    "佛山市": "佛山市",
    "顺德": "佛山市",
    "安顺": "安顺市",
    "安顺市": "安顺市",
    "太原": "太原市",
    "太原市": "太原市",
    "常德": "常德市",
    "常德市": "常德市",
    "桃源": "常德市",
    "桃源县": "常德市",
}

CUISINE_RULES = [
    ("面馆", "面食"),
    ("烧鸡", "烧鸡"),
    ("烧鹅", "烧味"),
    ("小炒", "小炒"),
    ("私房菜", "私房菜"),
    ("擂茶", "客家菜"),
    ("饭店", "地方菜"),
    ("餐馆", "家常菜"),
    ("酒店", "地方菜"),
    ("火锅", "火锅"),
    ("烤鱼", "烤鱼"),
    ("海鲜", "海鲜"),
    ("粤菜", "粤菜"),
    ("湘菜", "湘菜"),
    ("川菜", "川菜"),
]

TIP_RULES = {
    "booking": ["预约", "实名", "小程序", "公众号", "app", "电话预约", "联系客服", "提前约", "需组团"],
    "opening_hours": ["开放时间", "停止入场", "闭馆", "场次", "周一", "周二", "周三", "周四", "周五", "周六", "周日"],
    "ticket": ["门票", "票价", "购票", "免费", "元/人", "起", "收费"],
    "transport": ["地铁", "停车", "地址", "导航", "打车"],
    "warning": ["建议", "注意", "不要", "不接受散客", "提前到", "年龄太小", "关闭", "团队", "限流", "约不上"],
}

NOISY_FACT_KEYWORDS = {
    "看点蓝绝技",
    "集12生肖章",
    "谢谢宝",
    "想去",
    "重点参观",
    "高频标题包括",
    "小红书笔记",
}

SOCIAL_TITLE_NOISE = (
    "攻略",
    "打卡",
    "遛娃",
    "别再来了",
    "一日游",
    "合集",
    "后劲太大",
)

KNOWN_NOTE_POI_NAMES = (
    "三元牛奶工厂",
    "中国工艺美术馆",
    "中国考古博物馆",
    "北京二手淘货地图",
    "北京火箭总装厂",
    "北京珐琅厂",
    "北京红星二锅头酒厂",
    "北冰洋工厂",
    "小米汽车工厂",
    "爱慕时尚工厂",
)


def collapse_ws(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalize_name(value: object) -> str:
    text = collapse_ws(value)
    text = re.sub(r"[·•\s]", "", text)
    text = re.sub(r"（[^）]*）|\([^)]*\)", "", text)
    return text.lower()


def make_uuid(label: str) -> str:
    return str(uuid.uuid5(UUID_NAMESPACE, label))


def unique_rows_by_id(items: list[dict]) -> list[dict]:
    deduped: dict[str, dict] = {}
    for item in items:
        deduped[item["id"]] = item
    return list(deduped.values())


def parse_list_literal(value: object) -> list[str]:
    if value in (None, "", "None", "[]"):
        return []
    if isinstance(value, list):
        return [collapse_ws(item) for item in value if collapse_ws(item)]
    text = str(value)
    try:
        parsed = ast.literal_eval(text)
    except Exception:
        return []
    if not isinstance(parsed, list):
        return []
    return [collapse_ws(item) for item in parsed if collapse_ws(item)]


def parse_int(value: object) -> int:
    text = collapse_ws(value)
    if not text or text.lower() == "none":
        return 0
    match = re.search(r"-?\d+", text.replace(",", ""))
    return int(match.group(0)) if match else 0


def normalize_city(value: object) -> str | None:
    text = collapse_ws(value)
    if not text:
        return None
    if text in CITY_ALIASES:
        return CITY_ALIASES[text]
    if text.endswith("特别行政区"):
        return text
    if text.endswith("市") or text.endswith("自治州"):
        return text
    if text.endswith("地区"):
        return text
    if text in {"香港", "香港特别行政区"}:
        return "香港特别行政区"
    return CITY_ALIASES.get(text.replace("市", ""), f"{text}市")


def normalize_province(value: object) -> str | None:
    text = collapse_ws(value)
    return text or None


def split_lines(text: str) -> list[str]:
    lines = []
    for raw in re.split(r"[\n\r]+", text):
        line = collapse_ws(raw)
        if line:
            lines.append(line)
    return lines


def normalize_snippet(text: str) -> str:
    cleaned = collapse_ws(text)
    cleaned = re.sub(r"\[[^\]]+\]", "", cleaned)
    cleaned = re.sub(r"#[^#]+#", "", cleaned)
    cleaned = re.sub(r"[✨💡📍✅🐄✈️🛍️🥛🌱👶🥤🌾🥣🚗🏭⭐️🔥👉🏻👉🌟☀️🎫🧲📝⏰💰🚇🎟️😘]+", "", cleaned)
    cleaned = cleaned.replace("→", " ").replace("｜", " ").replace("|", " ")
    cleaned = re.sub(r"^[>＞\-—–:：、\.\s]+", "", cleaned)
    cleaned = collapse_ws(cleaned)
    return cleaned


def clean_fact_snippet(text: str) -> str:
    cleaned = normalize_snippet(text)
    cleaned = re.sub(r"^(\d+[\.、]|\(?[一二三四五六七八九十]+\)?[、\.])\s*", "", cleaned)
    cleaned = re.sub(r"^(开放时间|营业时间|时间|门票|票价|价格|预约|参观方式|入馆|交通|地址|地铁)[:：]?\s*", "", cleaned)
    cleaned = re.sub(r"【[^】]+】", "", cleaned)
    cleaned = collapse_ws(cleaned).strip("；;。,.，")
    return cleaned


def looks_like_valid_ticket(text: str) -> bool:
    if not text:
        return False
    if any(keyword in text for keyword in NOISY_FACT_KEYWORDS):
        return False
    return bool(re.search(r"(免费|收费|¥\s*\d+|\d+(?:\.\d+)?\s*(元|块|起))", text))


def normalize_ticket_text(text: str) -> str | None:
    cleaned = clean_fact_snippet(text)
    if not looks_like_valid_ticket(cleaned):
        return None
    if "免费" in cleaned:
        return "免费（需预约）" if "预约" in cleaned else "免费"
    price_match = re.search(r"(\d+(?:\.\d+)?)\s*(元|块|起)", cleaned)
    if not price_match:
        return None
    amount, unit = price_match.groups()
    if "团购" in cleaned:
        return f"团购票约{amount}{unit}"
    return f"{amount}{unit}"


def looks_like_opening_hours(text: str) -> bool:
    return bool(
        re.search(r"(\d{1,2}[:：]\d{2}|周[一二三四五六日天]|停止入场|闭馆|开放)", text)
    )


def normalize_opening_hours_text(text: str) -> str | None:
    cleaned = clean_fact_snippet(text)
    if not looks_like_opening_hours(cleaned):
        return None
    cleaned = re.sub(r"【[^】]+】", "", cleaned)
    cleaned = re.sub(r"^(时间是|开放时间为|营业时间为)\s*", "", cleaned)
    cleaned = collapse_ws(cleaned).strip("；;。,.，")
    if len(cleaned) > 80 and not re.search(r"\d{1,2}[:：]\d{2}", cleaned):
        return None
    if len(cleaned) > 40 and not re.search(r"\d{1,2}[:：]\d{2}", cleaned):
        return None
    return cleaned[:120] if cleaned else None


def normalize_address_text(text: str) -> str | None:
    cleaned = clean_fact_snippet(text)
    cleaned = re.split(r"(开放时间|营业时间|免费开放|停止入场|门票|票价)", cleaned)[0]
    cleaned = re.sub(r"（[^）]*(开放|闭馆|免费|时间)[^）]*）", "", cleaned)
    cleaned = re.sub(r"^(北京|北京市)[·•\-]\s*", "", cleaned)
    cleaned = cleaned.split("（")[0]
    cleaned = re.split(r"[，,；;]\s*(地铁|停车|导航)", cleaned)[0]
    cleaned = cleaned.replace("北京市政府", "").strip("；;。,.， ")
    cleaned = collapse_ws(cleaned)
    return cleaned or None


def build_default_attraction_summary(clean_name: str, category: str, tags: list[str]) -> str:
    if "博物馆" in clean_name or "美术馆" in clean_name or "艺术馆" in clean_name:
        return "专题博物馆，适合安排室内深度参观。"
    if "工厂" in clean_name or "酒厂" in clean_name or "总装厂" in clean_name or "工坊" in clean_name:
        return "工业参观点，可了解生产流程或参加主题体验。"
    if "地图" in clean_name or "集市" in clean_name or "淘货" in clean_name:
        return "主题路线整理，适合作为片区探索参考。"
    if category == "museum":
        return "适合第一次到访时安排的城市文化景点。"
    if category == "industrial_tourism":
        return "适合亲子或工业主题行程的参观点位。"
    if "shopping" in tags:
        return "适合步行闲逛和片区探索。"
    return "适合纳入城市探索行程的目的地。"


def extract_summary_line(clean_name: str, relevant_text: str) -> str | None:
    aliases = build_note_aliases(clean_name)
    for raw in split_lines(relevant_text):
        cleaned = clean_fact_snippet(raw)
        if not cleaned or len(cleaned) < 10 or len(cleaned) > 80:
            continue
        if any(keyword in cleaned for keyword in NOISY_FACT_KEYWORDS):
            continue
        if any(keyword in cleaned for keyword in ("预约", "开放时间", "门票", "票价", "地址", "地铁", "停车")):
            continue
        if any(keyword in cleaned for keyword in SOCIAL_TITLE_NOISE):
            continue
        if any(alias and alias in cleaned for alias in aliases):
            cleaned = cleaned.replace(clean_name, "", 1).strip("，,；;。 ")
        if cleaned:
            return cleaned[:80] + ("。" if not cleaned.endswith(("。", "！", "？")) else "")
    return None


def generate_embedding(project_url: str, anon_key: str, text: str, retries: int = 6) -> list[float]:
    payload = json.dumps({"text": text}).encode("utf-8")
    request = urllib.request.Request(
        f"{project_url}/functions/v1/generate-embedding",
        data=payload,
        headers={
            "Content-Type": "application/json",
            "apikey": anon_key,
            "Authorization": f"Bearer {anon_key}",
        },
        method="POST",
    )

    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                body = json.loads(response.read().decode("utf-8"))
            embedding = body.get("embedding")
            if not isinstance(embedding, list) or len(embedding) != 768:
                raise ValueError("embedding response is invalid")
            return embedding
        except Exception as exc:  # noqa: PERF203
            last_error = exc
            if attempt + 1 < retries:
                delay = 2.0 * (attempt + 1)
                if isinstance(exc, urllib.error.HTTPError) and exc.code == 429:
                    delay = 8.0 * (attempt + 1)
                time.sleep(delay)
            continue

    raise RuntimeError(f"generate-embedding failed after {retries} retries: {last_error}") from last_error


def generate_embeddings_batch(
    project_url: str,
    anon_key: str,
    texts: list[str],
    retries: int = 6,
) -> list[list[float]]:
    payload = json.dumps({"texts": texts}).encode("utf-8")
    request = urllib.request.Request(
        f"{project_url}/functions/v1/generate-embedding",
        data=payload,
        headers={
            "Content-Type": "application/json",
            "apikey": anon_key,
            "Authorization": f"Bearer {anon_key}",
        },
        method="POST",
    )

    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(request, timeout=180) as response:
                body = json.loads(response.read().decode("utf-8"))
            embeddings = body.get("embeddings")
            if not isinstance(embeddings, list) or len(embeddings) != len(texts):
                raise ValueError("embeddings response is invalid")
            if any(not isinstance(item, list) or len(item) != 768 for item in embeddings):
                raise ValueError("embeddings batch dimensions are invalid")
            return embeddings
        except Exception as exc:  # noqa: PERF203
            last_error = exc
            if attempt + 1 < retries:
                delay = 2.0 * (attempt + 1)
                if isinstance(exc, urllib.error.HTTPError) and exc.code == 429:
                    delay = 8.0 * (attempt + 1)
                time.sleep(delay)
            continue

    raise RuntimeError(f"generate-embedding batch failed after {retries} retries: {last_error}") from last_error


def attach_embeddings_batched(
    project_url: str,
    anon_key: str,
    items: list[dict],
    text_builder,
    batch_size: int,
    pause_seconds: float,
) -> None:
    if not items:
        return

    for start in range(0, len(items), batch_size):
        batch = items[start : start + batch_size]
        texts = [text_builder(item) for item in batch]
        embeddings = generate_embeddings_batch(project_url, anon_key, texts)
        for item, embedding in zip(batch, embeddings, strict=True):
            item["embedding"] = embedding
        if pause_seconds > 0 and start + batch_size < len(items):
            time.sleep(pause_seconds)


def vector_literal(values: list[float] | str | None) -> str | None:
    if not values:
        return None
    if isinstance(values, str):
        return values
    return "[" + ",".join(f"{value:.9f}".rstrip("0").rstrip(".") for value in values) + "]"


def post_json(url: str, anon_key: str, payload: dict) -> dict:
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url,
        data=data,
        headers={
            "Content-Type": "application/json",
            "apikey": anon_key,
            "Authorization": f"Bearer {anon_key}",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} calling {url}: {details}") from exc


def infer_cuisine(name: str, titles: list[str]) -> str | None:
    for keyword, cuisine in CUISINE_RULES:
        if keyword in name:
            return cuisine

    corpus = " ".join(titles)
    for keyword, cuisine in CUISINE_RULES:
        if keyword in corpus:
            return cuisine
    return None


def extract_specialties(titles: list[str]) -> list[str]:
    matches = []
    pattern = re.compile(r"([\u4e00-\u9fff]{2,12}(?:面|粉|饭|鸭|鸡|鹅|鱼|虾|蟹|肉|锅贴|锅|肠|饺|包|汤|糕|煲|粥|汽水|面包|烧麦|肠粉))")
    for title in titles:
        for match in pattern.findall(title):
            if match not in matches:
                matches.append(match)
            if len(matches) == 5:
                return matches
    return matches


def compute_position_popularity(media_count: int) -> float:
    raw = 0.32 + min(0.45, math.log(media_count + 1) / 4)
    return round(min(0.9, raw), 2)


def load_sheet_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    try:
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return []
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append({headers[index]: values[index] for index in range(len(headers))})
    return rows


def clean_attraction_name(filename_stem: str) -> str:
    name = filename_stem
    for suffix in ("参观", "手工"):
        if name.endswith(suffix):
            name = name[: -len(suffix)]
    return collapse_ws(name)


def infer_attraction_category(name: str) -> str:
    if "博物馆" in name or "美术馆" in name:
        return "博物馆"
    if "地图" in name or "淘货" in name:
        return "市场淘货"
    if "工厂" in name or "酒厂" in name or "总装厂" in name:
        return "工业参观"
    return "体验"


def infer_city_from_notes(name: str, note_rows: list[dict]) -> str:
    for alias, canonical in CITY_ALIASES.items():
        if alias and alias in name:
            return canonical

    corpus_parts = [name]
    for row in note_rows:
        corpus_parts.extend(
            [
                collapse_ws(row.get("标题")),
                collapse_ws(row.get("描述")),
                collapse_ws(row.get("ip归属地")),
            ]
        )
    corpus = " ".join(corpus_parts)

    counts = Counter()
    for alias, canonical in CITY_ALIASES.items():
        if alias and alias in corpus:
            counts[canonical] += corpus.count(alias)

    if counts:
        return counts.most_common(1)[0][0]
    return "北京市"


def build_note_aliases(clean_name: str) -> list[str]:
    aliases = {clean_name}
    aliases.add(clean_name.replace("工厂", ""))
    aliases.add(clean_name.replace("博物馆", ""))
    aliases.add(clean_name.replace("美术馆", ""))
    aliases.add(clean_name.replace("北京", ""))
    return [alias for alias in aliases if alias]


def mentions_foreign_poi(clean_name: str, text: str) -> bool:
    aliases = build_note_aliases(clean_name)
    cleaned = normalize_snippet(text)
    if any(alias and alias in cleaned for alias in aliases):
        return False

    for poi_name in KNOWN_NOTE_POI_NAMES:
        if poi_name == clean_name:
            continue
        if poi_name in cleaned or poi_name.replace("北京", "") in cleaned:
            return True

    return bool(
        re.search(
            r"[\u4e00-\u9fff]{2,16}(?:博物馆|美术馆|艺术馆|工厂|酒厂|总装厂|珐琅厂|淘货地图)",
            cleaned,
        )
    )


def extract_relevant_note_text(clean_name: str, note_rows: list[dict]) -> str:
    aliases = build_note_aliases(clean_name)
    blocks: list[str] = []

    for row in note_rows:
        description = collapse_ws(row.get("描述"))
        if not description:
            continue
        title = collapse_ws(row.get("标题"))
        lines = split_lines(str(row.get("描述") or ""))
        roundup_markers = sum(
            1
            for line in lines
            if re.match(r"^[✅✨💰💡📍🐄✈️🛍️🥛🌱👶🥤🌾🥣🚗🏭]|^【|^\d+[️⃣\.]", line)
        )
        title_has_alias = any(alias and alias in title for alias in aliases)
        matched = False
        for index, line in enumerate(lines):
            if any(alias and alias in line for alias in aliases):
                snippet = [line]
                for next_line in lines[index + 1 : index + 5]:
                    if any(alias and alias in next_line for alias in aliases):
                        break
                    if mentions_foreign_poi(clean_name, next_line):
                        break
                    snippet.append(next_line)
                blocks.append("\n".join(snippet))
                matched = True
        if not matched and roundup_markers < 4 and title_has_alias:
            blocks.append("\n".join(lines[:8]))

    return "\n\n".join(blocks)


def pick_first_line(text: str, keywords: list[str]) -> str | None:
    for line in split_lines(text):
        if any(keyword in line for keyword in keywords):
            return normalize_snippet(line)
    return None


def extract_address(clean_name: str, text: str) -> str | None:
    patterns = [
        r"(?:地址|📍|位置)[:：]\s*([^\n]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            raw = match.group(1)
            if mentions_foreign_poi(clean_name, raw):
                continue
            line = normalize_address_text(raw)
            if line:
                return line
    return None


def extract_ticket_price(clean_name: str, text: str) -> str | None:
    for raw in split_lines(text):
        if mentions_foreign_poi(clean_name, raw):
            continue
        if any(keyword in raw for keyword in ["门票", "票价", "免费", "购票", "预约"]):
            line = normalize_ticket_text(raw)
            if line:
                return line
    return None


def extract_opening_hours(clean_name: str, text: str) -> str | None:
    for raw in split_lines(text):
        if mentions_foreign_poi(clean_name, raw):
            continue
        if any(keyword in raw for keyword in ["开放时间", "停止入场", "闭馆", "场次", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]):
            line = normalize_opening_hours_text(raw)
            if line:
                return line
    return None


def extract_best_time(text: str) -> str | None:
    line = pick_first_line(text, ["周末", "工作日", "节假日", "旺季", "淡季"])
    if not line:
        return None
    return line[:120]


def extract_duration(text: str) -> str | None:
    values = []
    for match in re.finditer(r"(\d+(?:\.\d+)?)\s*(?:个?小时|小时)", text):
        hours = float(match.group(1))
        if 0.5 <= hours <= 8:
            values.append(hours)
    for match in re.finditer(r"(?:建议|至少|预留)?\s*(\d+(?:\.\d+)?)\s*小时", text):
        hours = float(match.group(1))
        if 0.5 <= hours <= 8:
            values.append(hours)
    if values:
        hours = round(sorted(values)[len(values) // 2], 1)
        return f"约{hours:g}小时"
    if "半天" in text:
        return "半天"
    return None


def build_note_tags(clean_name: str, note_rows: list[dict], category: str) -> list[str]:
    tags = []
    if category:
        tags.append(category)
    if "手工" in clean_name:
        tags.append("手作")
    if "工厂" in clean_name:
        tags.append("工厂")
    if "博物馆" in clean_name or "美术馆" in clean_name:
        tags.append("展览")

    for row in note_rows:
        for tag in parse_list_literal(row.get("标签")):
            if tag not in tags:
                tags.append(tag)
            if len(tags) >= 8:
                return tags
    return tags[:8]


def build_attraction_description(
    clean_name: str,
    city: str,
    category: str,
    note_rows: list[dict],
    comment_rows: list[dict],
    relevant_text: str,
    tags: list[str],
    opening_hours: str | None,
    ticket_price: str | None,
    recommended_duration: str | None,
) -> str:
    summary = build_default_attraction_summary(clean_name, category, tags).rstrip("。；; ")
    parts = [summary]
    parts.append(f"位于{city}")
    if opening_hours:
        parts.append(f"开放时间：{opening_hours}")
    if ticket_price:
        parts.append(f"门票：{ticket_price}")
    if recommended_duration:
        parts.append(f"建议预留{recommended_duration}")
    summary = "；".join(parts) + "。"
    return summary[:480]


def compute_note_popularity(note_rows: list[dict], comment_rows: list[dict]) -> float:
    likes = sum(parse_int(row.get("点赞数量")) for row in note_rows)
    saves = sum(parse_int(row.get("收藏数量")) for row in note_rows)
    comments = len(comment_rows) if comment_rows else sum(parse_int(row.get("评论数量")) for row in note_rows)
    raw = 0.35 + min(0.6, math.log(likes + saves + comments + 1) / 10)
    return round(min(0.95, raw), 2)


def extract_tip_candidates(texts: list[str]) -> dict[str, Counter]:
    buckets = {category: Counter() for category in TIP_RULES}
    for text in texts:
        for raw_line in split_lines(text):
            line = normalize_snippet(raw_line)
            if not line or len(line) < 8 or len(line) > 180:
                continue
            if line.endswith(("：", ":")):
                continue
            if line in {"预约方式", "步骤如下"}:
                continue
            for category, keywords in TIP_RULES.items():
                if any(keyword in line for keyword in keywords):
                    buckets[category][line] += 1
    return buckets


def build_travel_tips(clean_name: str, destination_id: str, relevant_text: str, comment_rows: list[dict]) -> list[dict]:
    texts = [relevant_text]
    texts.extend(collapse_ws(row.get("评论内容")) for row in comment_rows if collapse_ws(row.get("评论内容")))
    buckets = extract_tip_candidates(texts)
    tips = []
    title_map = {
        "booking": f"{clean_name}预约",
        "opening_hours": f"{clean_name}开放时间",
        "ticket": f"{clean_name}门票",
        "transport": f"{clean_name}交通",
        "warning": f"{clean_name}注意事项",
    }

    for category, counter in buckets.items():
        for content, _ in counter.most_common(3):
            tips.append(
                {
                    "id": make_uuid(f"tip:{destination_id}:{clean_name}:{category}:{content}"),
                    "destination_id": destination_id,
                    "category": category,
                    "title": title_map[category],
                    "content": content,
                    "is_important": category in {"booking", "opening_hours", "warning"},
                    "season_specific": None,
                }
            )
    return tips


def load_restaurants_from_workbook(position_path: Path):
    workbook = load_workbook(position_path, read_only=True)

    category_rows = workbook["Category"].iter_rows(min_row=2, values_only=True)
    category_map = {row[0]: collapse_ws(row[2]) for row in category_rows if row and row[0] is not None}

    platform_rows = workbook["Platform"].iter_rows(min_row=2, values_only=True)
    platform_map = {row[0]: collapse_ws(row[2]) for row in platform_rows if row and row[0] is not None}

    media_headers = next(workbook["Media"].iter_rows(min_row=1, max_row=1, values_only=True))
    media_idx = {header: index for index, header in enumerate(media_headers)}
    media_map = {}
    for row in workbook["Media"].iter_rows(min_row=2, values_only=True):
        media_id = row[media_idx["id"]]
        if not media_id:
            continue
        media_map[media_id] = {
            "id": media_id,
            "name": collapse_ws(row[media_idx["name"]]),
            "platform": platform_map.get(row[media_idx["platform_id"]]),
            "link_url": collapse_ws(row[media_idx["link_url"]]),
            "thumbnail": collapse_ws(row[media_idx["thumbnail"]]) or None,
            "upload_time": collapse_ws(row[media_idx["upload_time"]]),
        }

    media_by_position: dict[int, list[dict]] = defaultdict(list)
    for row in workbook["MediaPosition"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        media = media_map.get(row[2])
        position_id = row[3]
        if media and position_id:
            media_by_position[position_id].append(media)

    position_headers = next(workbook["Position"].iter_rows(min_row=1, max_row=1, values_only=True))
    position_idx = {header: index for index, header in enumerate(position_headers)}

    position_rows = [
        row
        for row in workbook["Position"].iter_rows(min_row=2, values_only=True)
        if collapse_ws(row[position_idx["name"]])
    ]

    base_id_counts: Counter[tuple[str, str]] = Counter()
    for row in position_rows:
        city = normalize_city(row[position_idx["city"]])
        if not city:
            continue
        name = collapse_ws(row[position_idx["name"]])
        base_id_counts[(city, name)] += 1

    restaurants = []
    seen_id_labels: set[str] = set()
    destination_stats: dict[str, dict] = defaultdict(
        lambda: {
            "restaurant_count": 0,
            "attraction_count": 0,
            "tip_count": 0,
            "media_count": 0,
            "province": None,
        }
    )

    for row in position_rows:
        name = collapse_ws(row[position_idx["name"]])
        if not name:
            continue
        city = normalize_city(row[position_idx["city"]])
        if not city:
            continue
        province = normalize_province(row[position_idx["province"]])
        district = collapse_ws(row[position_idx["district"]]) or None
        address = collapse_ws(row[position_idx["address"]]) or None
        media_items = sorted(
            media_by_position.get(row[position_idx["id"]], []),
            key=lambda item: item["upload_time"],
            reverse=True,
        )
        media_titles = [item["name"] for item in media_items if item["name"]]
        linked_platforms = []
        for item in media_items:
            platform = item.get("platform")
            if platform and platform not in linked_platforms:
                linked_platforms.append(platform)
        cuisine_type = infer_cuisine(name, media_titles)
        specialties = extract_specialties(media_titles)
        description_parts = [f"收录于大众点评位置库，城市为{city}"]
        if district:
            description_parts.append(f"区域：{district}")
        if address:
            description_parts.append(f"地址：{address}")
        if media_titles:
            description_parts.append(f"关联{len(media_titles)}条媒体内容")
            description_parts.append("相关标题：" + "、".join(media_titles[:2]))
        if linked_platforms:
            description_parts.append("内容平台：" + "、".join(linked_platforms[:2]))

        id_candidates = [f"restaurant:{city}:{name}"]
        if base_id_counts[(city, name)] > 1:
            if address:
                id_candidates.append(f"restaurant:{city}:{name}:{address}")
            lat = row[position_idx["latitude"]]
            lng = row[position_idx["longitude"]]
            if lat is not None and lng is not None:
                id_candidates.append(f"restaurant:{city}:{name}:{lat}:{lng}")
            id_candidates.append(f"restaurant:{city}:{name}:position:{row[position_idx['id']]}")

        chosen_id_label = id_candidates[0]
        if base_id_counts[(city, name)] > 1:
            for candidate in id_candidates[1:]:
                if candidate not in seen_id_labels:
                    chosen_id_label = candidate
                    break
        seen_id_labels.add(chosen_id_label)

        restaurants.append(
            {
                "id": make_uuid(chosen_id_label),
                "destination_id": make_uuid(f"destination:{city}"),
                "name": name,
                "name_en": None,
                "cuisine_type": cuisine_type,
                "description": "；".join(description_parts) + "。",
                "location_lat": row[position_idx["latitude"]],
                "location_lng": row[position_idx["longitude"]],
                "address": address,
                "price_range": None,
                "meal_type": category_map.get(row[position_idx["category_id"]]),
                "specialties": specialties or None,
                "dietary_options": None,
                "opening_hours": None,
                "reservation_required": False,
                "popularity_score": compute_position_popularity(len(media_items)),
                "avg_rating": None,
                "review_count": 0,
                "menu_image_url": next((item["thumbnail"] for item in media_items if item.get("thumbnail")), None),
                "city": city,
            }
        )

        stats = destination_stats[city]
        stats["restaurant_count"] += 1
        stats["media_count"] += len(media_items)
        if province and not stats["province"]:
            stats["province"] = province

    return restaurants, destination_stats


def load_attractions_from_notes(notes_dir: Path):
    attractions = []
    travel_tips = []
    destination_stats: dict[str, dict] = defaultdict(
        lambda: {
            "restaurant_count": 0,
            "attraction_count": 0,
            "tip_count": 0,
            "media_count": 0,
            "province": None,
        }
    )

    for path in sorted(notes_dir.glob("*.xlsx")):
        if path.name in SKIP_NOTE_FILES or path.name.endswith("_comments.xlsx"):
            continue

        note_rows = load_sheet_rows(path)
        if not note_rows:
            continue

        clean_name = clean_attraction_name(path.stem)
        category = infer_attraction_category(clean_name)
        city = infer_city_from_notes(clean_name, note_rows)
        destination_id = make_uuid(f"destination:{city}")

        comments_path = path.with_name(f"{path.stem}_comments.xlsx")
        comment_rows = load_sheet_rows(comments_path) if comments_path.exists() else []

        relevant_text = extract_relevant_note_text(clean_name, note_rows)
        opening_hours = extract_opening_hours(clean_name, relevant_text)
        ticket_price = extract_ticket_price(clean_name, relevant_text)
        recommended_duration = extract_duration(relevant_text)
        best_time_to_visit = extract_best_time(relevant_text)
        tags = build_note_tags(clean_name, note_rows, category)
        address = extract_address(clean_name, relevant_text)
        popularity_score = compute_note_popularity(note_rows, comment_rows)
        review_count = len(comment_rows) if comment_rows else sum(parse_int(row.get("评论数量")) for row in note_rows)
        description = build_attraction_description(
            clean_name,
            city,
            category,
            note_rows,
            comment_rows,
            relevant_text,
            tags,
            opening_hours,
            ticket_price,
            recommended_duration,
        )

        attraction = {
            "id": make_uuid(f"attraction:{city}:{clean_name}"),
            "destination_id": destination_id,
            "name": clean_name,
            "category": category,
            "description": description,
            "location_lat": None,
            "location_lng": None,
            "address": address,
            "opening_hours": {"text": opening_hours} if opening_hours else None,
            "ticket_price": ticket_price,
            "recommended_duration": recommended_duration,
            "best_time_to_visit": best_time_to_visit,
            "tags": tags,
            "popularity_score": popularity_score,
            "avg_rating": None,
            "review_count": review_count,
            "city": city,
        }
        attractions.append(attraction)

        tips = build_travel_tips(clean_name, destination_id, relevant_text, comment_rows)
        travel_tips.extend(tips)

        stats = destination_stats[city]
        stats["attraction_count"] += 1
        stats["tip_count"] += len(tips)

    return attractions, travel_tips, destination_stats


def build_destinations(destination_stats: dict[str, dict]) -> list[dict]:
    destinations = []
    for city in sorted(destination_stats):
        stats = destination_stats[city]
        parts = []
        if stats["restaurant_count"]:
            parts.append(f"收录{stats['restaurant_count']}家餐馆")
        if stats["media_count"]:
            parts.append(f"关联{stats['media_count']}条媒体内容")
        if stats["attraction_count"]:
            parts.append(f"沉淀{stats['attraction_count']}个景点/体验主题")
        if stats["tip_count"]:
            parts.append(f"提炼{stats['tip_count']}条出行提示")
        description = f"{city}目的地知识库" + ("，" + "，".join(parts) if parts else "") + "。"
        destinations.append(
            {
                "id": make_uuid(f"destination:{city}"),
                "name": city,
                "country": "中国",
                "description": description,
                "best_season": None,
                "average_budget_daily": None,
                "currency": "CNY",
                "timezone": "Asia/Shanghai",
                "province": stats.get("province"),
            }
        )
    return destinations


def merge_destination_stats(*all_stats: dict[str, dict]) -> dict[str, dict]:
    merged: dict[str, dict] = defaultdict(
        lambda: {
            "restaurant_count": 0,
            "attraction_count": 0,
            "tip_count": 0,
            "media_count": 0,
            "province": None,
        }
    )
    for stats in all_stats:
        for city, values in stats.items():
            target = merged[city]
            for key in ("restaurant_count", "attraction_count", "tip_count", "media_count"):
                target[key] += values.get(key, 0)
            if values.get("province") and not target.get("province"):
                target["province"] = values["province"]
    return merged


def upload_batches(
    project_url: str,
    anon_key: str,
    function_slug: str,
    destinations: list[dict],
    restaurants: list[dict],
    attractions: list[dict],
    travel_tips: list[dict],
    restaurant_batch_size: int,
) -> list[dict]:
    endpoint = f"{project_url}/functions/v1/{function_slug}"
    results = []

    destination_payload = [{**item, "embedding": vector_literal(item.get("embedding"))} for item in destinations]
    if destination_payload:
        results.append(
            post_json(
                endpoint,
                anon_key,
                {
                    "destinations": destination_payload,
                    "restaurants": [],
                    "dishes": [],
                    "attractions": [],
                    "travel_tips": [],
                    "replaceRestaurantDishes": False,
                },
            )
        )

    for start in range(0, len(restaurants), restaurant_batch_size):
        batch = restaurants[start : start + restaurant_batch_size]
        payload = [{**item, "embedding": vector_literal(item.get("embedding"))} for item in batch]
        results.append(
            post_json(
                endpoint,
                anon_key,
                {
                    "destinations": [],
                    "restaurants": payload,
                    "dishes": [],
                    "attractions": [],
                    "travel_tips": [],
                    "replaceRestaurantDishes": False,
                },
            )
        )

    if attractions or travel_tips:
        attraction_payload = [{**item, "embedding": vector_literal(item.get("embedding"))} for item in attractions]
        tip_payload = [{**item, "embedding": vector_literal(item.get("embedding"))} for item in travel_tips]
        results.append(
            post_json(
                endpoint,
                anon_key,
                {
                    "destinations": [],
                    "restaurants": [],
                    "dishes": [],
                    "attractions": attraction_payload,
                    "travel_tips": tip_payload,
                    "replaceRestaurantDishes": False,
                },
            )
        )

    return results


def main() -> int:
    parser = argparse.ArgumentParser(description="Import workbook restaurants/destinations and note-based attractions/travel tips.")
    parser.add_argument("--position-xlsx", default=DEFAULT_POSITION_XLSX)
    parser.add_argument("--notes-dir", default=DEFAULT_NOTES_DIR)
    parser.add_argument("--project-url", default=PROJECT_URL)
    parser.add_argument("--anon-key", default="")
    parser.add_argument("--function-slug", default=DEFAULT_FUNCTION_SLUG)
    parser.add_argument("--skip-embedding", action="store_true")
    parser.add_argument("--apply-via-function", action="store_true")
    parser.add_argument("--embedding-workers", type=int, default=8)
    parser.add_argument("--embedding-batch-size", type=int, default=20)
    parser.add_argument("--embedding-pause-seconds", type=float, default=0.4)
    parser.add_argument("--restaurant-batch-size", type=int, default=40)
    args = parser.parse_args()

    if not args.skip_embedding and not args.anon_key:
        parser.error("--anon-key is required unless --skip-embedding is used")
    if args.apply_via_function and not args.anon_key:
        parser.error("--anon-key is required when --apply-via-function is used")

    restaurants, restaurant_destination_stats = load_restaurants_from_workbook(Path(args.position_xlsx))
    attractions, travel_tips, note_destination_stats = load_attractions_from_notes(Path(args.notes_dir))
    destinations = unique_rows_by_id(build_destinations(merge_destination_stats(restaurant_destination_stats, note_destination_stats)))
    restaurants = unique_rows_by_id(restaurants)
    attractions = unique_rows_by_id(attractions)
    travel_tips = unique_rows_by_id(travel_tips)

    if not args.skip_embedding:
        attach_embeddings_batched(
            args.project_url,
            args.anon_key,
            destinations,
            lambda item: " ".join(part for part in [item["name"], item["description"], item["country"]] if part),
            max(1, args.embedding_batch_size),
            max(0.0, args.embedding_pause_seconds),
        )
        attach_embeddings_batched(
            args.project_url,
            args.anon_key,
            restaurants,
            lambda item: " ".join(
                part
                for part in [
                    item.get("city") or "",
                    item["name"],
                    item.get("cuisine_type") or "",
                    item.get("description") or "",
                    " ".join(item.get("specialties") or []),
                    item.get("address") or "",
                ]
                if part
            ),
            max(1, args.embedding_batch_size),
            max(0.0, args.embedding_pause_seconds),
        )
        attach_embeddings_batched(
            args.project_url,
            args.anon_key,
            attractions,
            lambda item: " ".join(
                part
                for part in [
                    item.get("city") or "",
                    item["name"],
                    item.get("category") or "",
                    item.get("description") or "",
                    " ".join(item.get("tags") or []),
                    item.get("ticket_price") or "",
                    item.get("recommended_duration") or "",
                    item.get("address") or "",
                ]
                if part
            ),
            max(1, args.embedding_batch_size),
            max(0.0, args.embedding_pause_seconds),
        )
        attach_embeddings_batched(
            args.project_url,
            args.anon_key,
            travel_tips,
            lambda item: " ".join(
                part
                for part in [
                    item.get("category") or "",
                    item.get("title") or "",
                    item.get("content") or "",
                ]
                if part
            ),
            max(1, args.embedding_batch_size),
            max(0.0, args.embedding_pause_seconds),
        )

    upload_results = []
    if args.apply_via_function:
        upload_results = upload_batches(
            args.project_url,
            args.anon_key,
            args.function_slug,
            destinations,
            restaurants,
            attractions,
            travel_tips,
            max(1, args.restaurant_batch_size),
        )

    summary = {
        "destinations": len(destinations),
        "restaurants": len(restaurants),
        "attractions": len(attractions),
        "travel_tips": len(travel_tips),
        "restaurant_cities": len({item["city"] for item in restaurants}),
        "restaurants_with_address": sum(1 for item in restaurants if item.get("address")),
        "attractions_with_address": sum(1 for item in attractions if item.get("address")),
        "attractions_with_opening_hours": sum(1 for item in attractions if item.get("opening_hours")),
        "upload_batches": len(upload_results),
        "upload_results": upload_results,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
